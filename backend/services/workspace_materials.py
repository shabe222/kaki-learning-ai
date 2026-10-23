"""Persistent upload, parsing, chunking, binding, and retrieval for workspace materials."""
from __future__ import annotations

import re
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from database.models import (
    WorkspaceConversationMaterial,
    WorkspaceMaterial,
    WorkspaceMaterialChunk,
)
from services.workspace_input_adapters import InputAdapterError, ocr_image_bytes


MAX_MATERIAL_BYTES = 30 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_DOCX_ENTRIES = 4000
MAX_PDF_PAGES = 500
CHUNK_CHARS = 1600
CHUNK_OVERLAP = 140
MAX_CONTEXT_CHARS = 7600
MAX_CATEGORY_CHARS = 60
MAX_TAGS = 12
MAX_TAG_CHARS = 32

MIME_BY_SUFFIX = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
}


class MaterialError(ValueError):
    pass


class MaterialTypeError(MaterialError):
    pass


class MaterialTooLargeError(MaterialError):
    pass


class MaterialAccessError(MaterialError):
    pass


@dataclass(frozen=True)
class ParsedSection:
    heading: str
    text: str
    page_start: Optional[int] = None
    page_end: Optional[int] = None
    level: int = 1
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ParsedMaterial:
    file_type: str
    sections: List[ParsedSection]
    page_count: Optional[int]
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class MaterialContext:
    material_ids: List[str]
    text: str
    citations: List[Dict[str, Any]]


MATERIAL_REFERENCE_PATTERN = re.compile(
    r"(?:这|那|刚才|之前|上面)(?:份|个|张|些)?(?:文件|资料|文档|图片|教材|PDF)"
    r"|(?:上传|附件|文件|资料|文档|图片|教材|PDF|DOCX|Markdown)"
    r"|第[一二三四五六七八九十百零\d]+(?:章|节|页|份)"
    r"|(?:只|仅)(?:使用|根据|依据|参考)",
    re.I,
)


def should_use_material_context(message: str, policy: str = "auto") -> bool:
    """Decide whether active attachments should be retrieved for this turn."""
    if policy == "exclude":
        return False
    if policy == "only":
        return True
    return bool(MATERIAL_REFERENCE_PATTERN.search(str(message or "")))


def should_retrieve_material_context(
    message: str,
    policy: str = "auto",
    requested_material_ids: Optional[Sequence[str]] = None,
) -> bool:
    """Resolve the final retrieval gate; explicit exclude always wins over stale IDs."""
    if policy == "exclude":
        return False
    return bool(requested_material_ids) or should_use_material_context(message, policy)


def _clean_file_name(file_name: str) -> str:
    cleaned = Path(str(file_name or "").replace("\x00", "")).name.strip()
    if not cleaned or len(cleaned) > 255:
        raise MaterialError("文件名无效。")
    return cleaned


def normalize_material_category(value: Optional[str]) -> str:
    category = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(category) > MAX_CATEGORY_CHARS:
        raise MaterialError(f"分类不能超过 {MAX_CATEGORY_CHARS} 个字符。")
    return category or "未分类"


def normalize_material_tags(values: Optional[Sequence[str]]) -> List[str]:
    if values is None:
        return []
    if isinstance(values, (str, bytes)):
        raise MaterialError("标签必须是字符串数组。")
    normalized: List[str] = []
    seen = set()
    for raw in values:
        if not isinstance(raw, str):
            raise MaterialError("标签必须是字符串数组。")
        tag = re.sub(r"\s+", " ", raw).strip()
        if not tag:
            continue
        if len(tag) > MAX_TAG_CHARS:
            raise MaterialError(f"单个标签不能超过 {MAX_TAG_CHARS} 个字符。")
        key = tag.casefold()
        if key in seen:
            continue
        seen.add(key)
        normalized.append(tag)
        if len(normalized) > MAX_TAGS:
            raise MaterialError(f"标签不能超过 {MAX_TAGS} 个。")
    return normalized


def _file_type(file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return "image"
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    if suffix == ".txt":
        return "txt"
    if suffix in {".md", ".markdown"}:
        return "markdown"
    raise MaterialTypeError("文件格式不支持，请上传图片、PDF、DOCX、TXT 或 Markdown。")


def validate_material_payload(payload: bytes, file_name: str) -> None:
    if not payload:
        raise MaterialError("文件内容为空。")
    if len(payload) > MAX_MATERIAL_BYTES:
        raise MaterialTooLargeError("文件不能超过 30 MB。")
    _file_type(_clean_file_name(file_name))


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = payload.decode(encoding)
            if text.strip():
                return text
        except UnicodeDecodeError:
            continue
    raise MaterialError("文本文件编码无法识别，请使用 UTF-8。")


def _plain_sections(text: str, markdown: bool) -> List[ParsedSection]:
    if not markdown:
        return [ParsedSection(heading="正文", text=text.strip())]
    sections: List[ParsedSection] = []
    heading = "正文"
    level = 1
    buffer: List[str] = []
    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if match:
            if "\n".join(buffer).strip():
                sections.append(ParsedSection(heading=heading, text="\n".join(buffer).strip(), level=level))
            heading = match.group(2).strip()
            level = len(match.group(1))
            buffer = []
        else:
            buffer.append(line)
    if "\n".join(buffer).strip() or not sections:
        sections.append(ParsedSection(heading=heading, text="\n".join(buffer).strip(), level=level))
    return sections


def _parse_pdf(payload: bytes) -> ParsedMaterial:
    try:
        import pdfplumber
        with pdfplumber.open(BytesIO(payload)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise MaterialError("PDF 页数不能超过 500 页。")
            extracted = [str(page.extract_text() or "").strip() for page in pdf.pages]
    except MaterialError:
        raise
    except Exception as exc:
        raise MaterialError("PDF 无法读取，请确认文件没有损坏或加密。") from exc

    ocr_pages: List[int] = []
    if any(not text for text in extracted):
        try:
            import fitz
            document = fitz.open(stream=payload, filetype="pdf")
            for index, text in enumerate(extracted):
                if text:
                    continue
                page = document.load_page(index)
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
                result = ocr_image_bytes(pixmap.tobytes("png"), content_type="image/png")
                extracted[index] = result.text
                ocr_pages.append(index + 1)
            document.close()
        except (InputAdapterError, Exception) as exc:
            if not any(extracted):
                raise MaterialError("扫描 PDF OCR 失败，请上传更清晰的文件。") from exc

    sections = []
    for index, text in enumerate(extracted):
        if not text:
            continue
        first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")
        sections.append(ParsedSection(
            heading=first_line[:160] or f"第 {index + 1} 页",
            text=text,
            page_start=index + 1,
            page_end=index + 1,
            metadata={"ocr": index + 1 in ocr_pages},
        ))
    if not sections:
        raise MaterialError("PDF 中没有可提取的文字。")
    return ParsedMaterial("pdf", sections, len(extracted), {"ocr_pages": ocr_pages})


def _validate_docx_archive(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_ENTRIES:
                raise MaterialError("DOCX 内容过于复杂。")
            if sum(info.file_size for info in infos) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise MaterialError("DOCX 解压后内容过大。")
            if "word/document.xml" not in {info.filename for info in infos}:
                raise MaterialError("DOCX 缺少正文结构。")
    except MaterialError:
        raise
    except zipfile.BadZipFile as exc:
        raise MaterialError("DOCX 无法读取，请确认文件没有损坏。") from exc


def _parse_docx(payload: bytes) -> ParsedMaterial:
    _validate_docx_archive(payload)
    try:
        from docx import Document
        document = Document(BytesIO(payload))
    except Exception as exc:
        raise MaterialError("DOCX 无法解析。") from exc

    sections: List[ParsedSection] = []
    heading = "正文"
    level = 1
    buffer: List[str] = []

    def flush() -> None:
        nonlocal buffer
        text = "\n".join(buffer).strip()
        if text:
            sections.append(ParsedSection(heading=heading, text=text, level=level))
        buffer = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = str(paragraph.style.name or "") if paragraph.style else ""
        match = re.search(r"Heading\s+(\d+)", style_name, re.I)
        if match:
            flush()
            heading = text
            level = min(int(match.group(1)), 6)
        else:
            buffer.append(text)
    flush()

    table_lines: List[str] = []
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(values):
                table_lines.append(" | ".join(values))
    if table_lines:
        if sections:
            last = sections[-1]
            sections[-1] = ParsedSection(
                heading=last.heading,
                text=(last.text + "\n\n表格：\n" + "\n".join(table_lines)).strip(),
                level=last.level,
                metadata={"has_tables": True},
            )
        else:
            sections.append(ParsedSection("表格", "\n".join(table_lines)))
    if not sections:
        raise MaterialError("DOCX 中没有可提取的文字。")
    return ParsedMaterial("docx", sections, None, {"table_count": len(document.tables)})


def parse_material_bytes(payload: bytes, file_name: str, content_type: str) -> ParsedMaterial:
    clean_name = _clean_file_name(file_name)
    validate_material_payload(payload, clean_name)
    kind = _file_type(clean_name)
    if kind == "image":
        result = ocr_image_bytes(payload, content_type=MIME_BY_SUFFIX[Path(clean_name).suffix.lower()])
        return ParsedMaterial("image", [ParsedSection("图片文字", result.text, 1, 1)], 1, {
            "ocr_engine": result.engine,
            "ocr_confidence": result.confidence,
            "line_count": result.line_count,
        })
    if kind == "pdf":
        return _parse_pdf(payload)
    if kind == "docx":
        return _parse_docx(payload)
    text = _decode_text(payload)
    return ParsedMaterial(kind, _plain_sections(text, markdown=kind == "markdown"), None, {})


def _chunk_sections(sections: Sequence[ParsedSection]) -> List[Dict[str, Any]]:
    chunks: List[Dict[str, Any]] = []
    for section in sections:
        text = section.text.strip()
        start = 0
        while start < len(text):
            end = min(len(text), start + CHUNK_CHARS)
            if end < len(text):
                boundary = max(text.rfind("\n", start, end), text.rfind("。", start, end))
                if boundary > start + 500:
                    end = boundary + 1
            part = text[start:end].strip()
            if part:
                chunks.append({
                    "heading": section.heading[:300],
                    "page_start": section.page_start,
                    "page_end": section.page_end,
                    "content": part,
                    "source_metadata": dict(section.metadata or {}),
                })
            if end >= len(text):
                break
            start = max(start + 1, end - CHUNK_OVERLAP)
    return chunks


async def create_workspace_material(
    db: AsyncSession,
    *,
    user_id: int,
    thread_id: str,
    file_name: str,
    content_type: str,
    payload: bytes,
    storage_root: Optional[Path] = None,
) -> WorkspaceMaterial:
    clean_name = _clean_file_name(file_name)
    validate_material_payload(payload, clean_name)
    if not thread_id or len(thread_id) > 100:
        raise MaterialError("会话 ID 无效。")
    material_id = "material_" + str(uuid.uuid4())
    suffix = Path(clean_name).suffix.lower()
    root = Path(storage_root) if storage_root else Path(__file__).resolve().parents[1] / "uploads" / "workspace_materials"
    user_root = root / str(user_id)
    user_root.mkdir(parents=True, exist_ok=True)
    final_path = user_root / f"{material_id}{suffix}"
    temporary_path = user_root / f".{material_id}.uploading"
    temporary_path.write_bytes(payload)
    temporary_path.replace(final_path)

    material = WorkspaceMaterial(
        material_id=material_id,
        user_id=user_id,
        file_name=clean_name,
        file_type=_file_type(clean_name),
        mime_type=MIME_BY_SUFFIX[suffix],
        size_bytes=len(payload),
        storage_path=str(final_path),
        parse_status="parsing",
        structure=[],
        parse_metadata={},
    )
    db.add(material)
    # The binding references a stable string key rather than the ORM primary key.
    # Flush the material first so SQLite/Postgres FK enforcement never observes
    # the binding before its parent row.
    await db.flush()
    binding = WorkspaceConversationMaterial(
        user_id=user_id, thread_id=thread_id, material_id=material_id, is_active=True,
    )
    db.add(binding)
    await db.commit()

    try:
        parsed = parse_material_bytes(payload, clean_name, content_type)
        chunks = _chunk_sections(parsed.sections)
        if not chunks:
            raise MaterialError("文件中没有可检索的文字。")
        for index, chunk in enumerate(chunks):
            db.add(WorkspaceMaterialChunk(
                material_id=material_id,
                user_id=user_id,
                chunk_index=index,
                heading=chunk["heading"],
                page_start=chunk["page_start"],
                page_end=chunk["page_end"],
                content=chunk["content"],
                char_count=len(chunk["content"]),
                source_metadata=chunk["source_metadata"],
            ))
        material.parse_status = "ready"
        material.parse_error = None
        material.page_count = parsed.page_count
        material.section_count = len(parsed.sections)
        material.char_count = sum(len(section.text) for section in parsed.sections)
        material.structure = [{
            "heading": section.heading,
            "level": section.level,
            "page_start": section.page_start,
            "page_end": section.page_end,
        } for section in parsed.sections]
        material.parse_metadata = {
            **parsed.metadata,
            "preview_text": "\n".join(section.text for section in parsed.sections)[:2400],
        }
        await db.commit()
        await db.refresh(material)
        return material
    except Exception as exc:
        material.parse_status = "failed"
        material.parse_error = str(exc)[:500]
        await db.commit()
        if isinstance(exc, MaterialError):
            raise
        if isinstance(exc, InputAdapterError):
            raise MaterialError(str(exc)) from exc
        raise MaterialError("文件解析失败。") from exc


async def list_conversation_materials(db: AsyncSession, user_id: int, thread_id: str) -> List[WorkspaceMaterial]:
    result = await db.execute(
        select(WorkspaceMaterial)
        .join(WorkspaceConversationMaterial, WorkspaceConversationMaterial.material_id == WorkspaceMaterial.material_id)
        .where(
            WorkspaceConversationMaterial.user_id == user_id,
            WorkspaceConversationMaterial.thread_id == thread_id,
            WorkspaceConversationMaterial.is_active.is_(True),
            WorkspaceMaterial.user_id == user_id,
        )
        .order_by(WorkspaceConversationMaterial.attached_at.desc())
    )
    return list(result.scalars().all())


async def list_library_materials(
    db: AsyncSession,
    user_id: int,
    *,
    query: Optional[str] = None,
    file_type: Optional[str] = None,
    category: Optional[str] = None,
) -> List[WorkspaceMaterial]:
    result = await db.execute(
        select(WorkspaceMaterial)
        .where(WorkspaceMaterial.user_id == user_id)
        .order_by(WorkspaceMaterial.created_at.desc())
    )
    items = list(result.scalars().all())
    normalized_type = str(file_type or "").strip().lower()
    normalized_category = normalize_material_category(category) if category else ""
    normalized_query = re.sub(r"\s+", " ", str(query or "")).strip().casefold()
    if normalized_type:
        items = [item for item in items if str(item.file_type or "").lower() == normalized_type]
    if normalized_category:
        items = [
            item for item in items
            if normalize_material_category((item.parse_metadata or {}).get("category")) == normalized_category
        ]
    if normalized_query:
        filtered: List[WorkspaceMaterial] = []
        for item in items:
            metadata = dict(item.parse_metadata or {})
            haystack = "\n".join([
                str(item.file_name or ""),
                normalize_material_category(metadata.get("category")),
                " ".join(normalize_material_tags(metadata.get("tags") or [])),
                str(metadata.get("preview_text") or ""),
            ]).casefold()
            if normalized_query in haystack:
                filtered.append(item)
        items = filtered
    return items


async def detach_workspace_material(db: AsyncSession, user_id: int, thread_id: str, material_id: str) -> bool:
    result = await db.execute(select(WorkspaceConversationMaterial).where(
        WorkspaceConversationMaterial.user_id == user_id,
        WorkspaceConversationMaterial.thread_id == thread_id,
        WorkspaceConversationMaterial.material_id == material_id,
        WorkspaceConversationMaterial.is_active.is_(True),
    ))
    binding = result.scalar_one_or_none()
    if binding is None:
        return False
    binding.is_active = False
    await db.commit()
    return True


async def attach_workspace_material(db: AsyncSession, user_id: int, thread_id: str, material_id: str) -> bool:
    if not thread_id or len(thread_id) > 100:
        raise MaterialError("会话 ID 无效。")
    material_result = await db.execute(select(WorkspaceMaterial).where(
        WorkspaceMaterial.user_id == user_id,
        WorkspaceMaterial.material_id == material_id,
    ))
    if material_result.scalar_one_or_none() is None:
        return False
    binding_result = await db.execute(select(WorkspaceConversationMaterial).where(
        WorkspaceConversationMaterial.user_id == user_id,
        WorkspaceConversationMaterial.thread_id == thread_id,
        WorkspaceConversationMaterial.material_id == material_id,
    ))
    binding = binding_result.scalar_one_or_none()
    if binding is None:
        db.add(WorkspaceConversationMaterial(
            user_id=user_id,
            thread_id=thread_id,
            material_id=material_id,
            is_active=True,
        ))
    else:
        binding.is_active = True
        binding.attached_at = datetime.utcnow()
    await db.commit()
    return True


def _query_terms(query: str) -> List[str]:
    normalized = str(query or "").lower()
    terms = re.findall(r"[a-z0-9_-]{2,}|[\u4e00-\u9fff]{2,}", normalized)
    expanded: List[str] = []
    for term in terms:
        expanded.append(term)
        if re.fullmatch(r"[\u4e00-\u9fff]+", term) and len(term) > 4:
            expanded.extend(term[index:index + 2] for index in range(len(term) - 1))
    return list(dict.fromkeys(expanded))


async def build_conversation_material_context(
    db: AsyncSession,
    *,
    user_id: int,
    thread_id: str,
    query: str,
    requested_material_ids: Optional[Sequence[str]] = None,
    scope: Optional[str] = None,
) -> MaterialContext:
    available = await list_conversation_materials(db, user_id, thread_id)
    by_id = {item.material_id: item for item in available}
    if requested_material_ids:
        missing = [item for item in requested_material_ids if item not in by_id]
        if missing:
            raise MaterialAccessError("资料不存在或未绑定当前会话。")
        selected = [by_id[item] for item in requested_material_ids]
    else:
        selected = available
    selected = [item for item in selected if item.parse_status == "ready"]
    if not selected:
        return MaterialContext([], "", [])

    named = [item for item in selected if item.file_name.lower() in query.lower()]
    if named:
        selected = named
    elif len(selected) > 1 and not any(term in query for term in ("两份", "这些", "全部", "比较", "对比")):
        selected = selected[:1]

    ids = [item.material_id for item in selected]
    result = await db.execute(select(WorkspaceMaterialChunk).where(
        WorkspaceMaterialChunk.user_id == user_id,
        WorkspaceMaterialChunk.material_id.in_(ids),
    ))
    chunks = list(result.scalars().all())
    terms = _query_terms((scope or "") + " " + query)

    def score(chunk: WorkspaceMaterialChunk) -> int:
        haystack = ((chunk.heading or "") + "\n" + chunk.content).lower()
        return sum((3 if term in (chunk.heading or "").lower() else 1) for term in terms if term in haystack)

    chunks.sort(key=lambda item: (-score(item), item.material_id, item.chunk_index))
    chosen = chunks[:6]
    parts: List[str] = []
    citations: List[Dict[str, Any]] = []
    used_chars = 0
    for chunk in chosen:
        material = by_id[chunk.material_id]
        location = ""
        if chunk.page_start:
            location = f"第 {chunk.page_start} 页"
        if chunk.heading:
            location = (location + " · " if location else "") + chunk.heading
        header = f"【资料：{material.file_name}" + (f"｜{location}" if location else "") + "】"
        block = header + "\n" + chunk.content.strip()
        remaining = MAX_CONTEXT_CHARS - used_chars
        if remaining <= 120:
            break
        if len(block) > remaining:
            block = block[:remaining] + "…"
        parts.append(block)
        used_chars += len(block)
        citations.append({
            "material_id": chunk.material_id,
            "file_name": material.file_name,
            "heading": chunk.heading,
            "page_start": chunk.page_start,
            "page_end": chunk.page_end,
            "chunk_index": chunk.chunk_index,
        })
    now = datetime.utcnow()
    await db.execute(
        WorkspaceConversationMaterial.__table__.update()
        .where(
            WorkspaceConversationMaterial.user_id == user_id,
            WorkspaceConversationMaterial.thread_id == thread_id,
            WorkspaceConversationMaterial.material_id.in_(ids),
        )
        .values(last_referenced_at=now)
    )
    await db.commit()
    return MaterialContext(ids, "\n\n".join(parts), citations)


async def delete_library_material(db: AsyncSession, user_id: int, material_id: str) -> bool:
    result = await db.execute(select(WorkspaceMaterial).where(
        WorkspaceMaterial.user_id == user_id,
        WorkspaceMaterial.material_id == material_id,
    ))
    material = result.scalar_one_or_none()
    if material is None:
        return False
    await db.execute(delete(WorkspaceConversationMaterial).where(WorkspaceConversationMaterial.material_id == material_id))
    await db.execute(delete(WorkspaceMaterialChunk).where(WorkspaceMaterialChunk.material_id == material_id))
    await db.delete(material)
    await db.commit()
    try:
        Path(material.storage_path).unlink(missing_ok=True)
    except OSError:
        pass
    return True
