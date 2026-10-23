"""Real slide generation and PPTX export for the standalone AI workspace.

Dual-theme contract: every presentation carries a version-level theme snapshot
(`presentation_theme` = kaki | ocean) captured at generation time. Old artifacts
keep their snapshot; switching the product theme later never mutates them.
Both themes are strict Morandi: low saturation, mid-low contrast, no pure
white/black, no neon, no e-commerce orange, no tech blue.
"""
from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any, Mapping

from config import settings
from services.generation_draft import GenerationDraftBlocked, presentation_draft
from services.skill_library import skill_prompt_prefix
from services.visual_theme import (
    PRESENTATION_THEME_SNAPSHOT,
    normalize_presentation_theme,
    presentation_theme_tokens,
)


def _text(value: Any, limit: int = 3000) -> str:
    return str(value or "").strip()[:limit]


def _mapping(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def _extract_json(raw: Any) -> dict:
    text = _text(raw, 40000).replace("\ufeff", "")
    if text.startswith("```"):
        text = text.split("\n", 1)[-1].rsplit("```", 1)[0].strip()
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("presentation agent returned no JSON object")
    value = json.loads(text[start:end + 1])
    if not isinstance(value, dict):
        raise ValueError("presentation agent returned invalid JSON")
    return value


# 双主题快照契约（版本级；生成时定格，之后不随产品主题变化）的令牌本体在
# services.visual_theme，与视频逐帧渲染共用，避免两个通道各留一份调色。



def normalize_slide_deck(value: Mapping[str, Any], *, include_interactions: bool) -> dict:
    title = _text(value.get("title"), 200) or "AI 学习演示"
    source_slides = value.get("slides")
    if isinstance(source_slides, dict):
        source_slides = source_slides.get("slides")
    if not isinstance(source_slides, list) or not source_slides:
        raise ValueError("presentation agent returned no slides")
    slides = []
    for index, raw in enumerate(source_slides[:30]):
        if not isinstance(raw, Mapping):
            continue
        kind = _text(raw.get("type"), 30) or "content"
        slide = {
            "index": index,
            "type": kind,
            "title": _text(raw.get("title"), 200) or (title if index == 0 else f"第 {index + 1} 页"),
            "subtitle": _text(raw.get("subtitle"), 300),
            "body": _text(raw.get("body"), 4000),
            "items": [_text(item, 400) for item in (raw.get("items") or raw.get("points") or []) if _text(item, 400)],
            "code": _text(raw.get("code"), 5000),
            "language": _text(raw.get("lang") or raw.get("language"), 30),
            "narration": _text(raw.get("narration"), 5000),
            "illustration": _text(raw.get("illustration"), 300),
            "image": _text(raw.get("image"), 600),
            "design": {
                "composition": _text(_mapping(raw.get("design")).get("composition"), 240),
                "image_position": (
                    str(_mapping(raw.get("design")).get("image_position") or "inline").lower()
                    if str(_mapping(raw.get("design")).get("image_position") or "inline").lower()
                    in {"right", "left", "background", "full_bleed", "inline", "none"}
                    else "inline"
                ),
                "density": (
                    str(_mapping(raw.get("design")).get("density") or "balanced").lower()
                    if str(_mapping(raw.get("design")).get("density") or "balanced").lower()
                    in {"airy", "balanced", "dense"} else "balanced"
                ),
                "emphasis": _text(_mapping(raw.get("design")).get("emphasis"), 180),
            },
        }
        question = _text(raw.get("question"), 500)
        options = [_text(item, 300) for item in (raw.get("options") or []) if _text(item, 300)]
        if include_interactions and question and len(options) >= 2:
            answer = raw.get("answer", raw.get("correct_index", 0))
            try:
                answer = int(answer)
            except (TypeError, ValueError):
                answer = 0
            if not 0 <= answer < len(options):
                answer = 0
            slide["interaction"] = {
                "question": question,
                "options": options,
                "correct_index": answer,
                "explanation": _text(raw.get("explanation"), 1000),
            }
        slides.append(slide)
    if not slides:
        raise ValueError("presentation agent returned no usable slides")
    return {"title": title, "slides": slides}


async def generate_workspace_presentation(
    *, topic: str, delivery_mode: str, learning_goal: str = "", level: str = "",
    duration_minutes: int | None = None, slide_count: int | None = None,
    presentation_style: str = "",
    source_context: str = "",
    presentation_theme: str = "kaki",
    learner_context: str = "",
) -> dict:
    """Generate real slide data through the configured production LLM; never mock."""
    include_interactions = delivery_mode == "interactive_lesson"
    slide_hint = int(slide_count) if slide_count else max(5, min(16, int((duration_minutes or 6) * 1.2)))
    slide_requirement = (
        f"必须恰好 {slide_hint} 页"
        if slide_count else f"建议约 {slide_hint} 页"
    )
    interaction_rule = (
        "至少包含 1 页 quiz，每道互动题给出 options、answer(从0开始)和 explanation。"
        if include_interactions else
        "不要生成测验页，不要包含答案字段；这是供用户查看和导出的普通 PPT。"
    )
    theme_key = normalize_presentation_theme(presentation_theme)
    # 没有现成来源资料时，先为这个主题取一次检索来源。任意主题的内容因此有依据可挂，
    # 而不是全靠模型记忆 —— 闸门能判"像不像专家"，判不了"对不对"。
    from services.generation_grounding import grounding_for_topic, grounding_prompt_rule

    grounding: dict[str, Any] = {"searched": False, "source_context": "", "source_refs": [], "note": ""}
    if not str(source_context or "").strip():
        grounding = await grounding_for_topic(topic, learning_goal=learning_goal)
        source_context = str(grounding.get("source_context") or "")
    theme_mood = PRESENTATION_THEME_SNAPSHOT[theme_key]["mood"]
    theme_hint = f"这份 PPT 将使用「{theme_key}」主题的视觉令牌（{theme_mood}）。主题只影响配色和材质氛围，绝不替代你对页面构图、例子、图文关系和信息节奏的设计判断。"
    prompt = f"""{skill_prompt_prefix("resource-art-direction", "presentation-authoring", max_chars_per_skill=1600)}你是 PPT 与互动课堂内容生成智能体。请生成严格 JSON，不要输出 Markdown 代码围栏。
主题：{topic}
交付模式：{delivery_mode}
学习目标：{learning_goal or '按主题组织核心内容'}
学习水平：{level or '通用'}
预计时长：{duration_minutes or 6} 分钟
呈现风格：{presentation_style or '清晰、简洁、适合学习'}
页数要求：{slide_requirement}；封面和收束页也计入总页数，不得额外增加页面。
{theme_hint}
学习者与本轮任务依据：{learner_context[:3000] if learner_context else '暂无额外背景'}
来源资料：{source_context[:6000] if source_context else '无'}
来源规则：{grounding_prompt_rule(grounding)}

输出结构：
{{"title":"标题","slides":[{{"type":"title|goals|content|code|quiz|summary","title":"页标题","subtitle":"可选","body":"可选正文","items":["要点"],"code":"可选代码","lang":"可选语言","narration":"本页自然讲解","illustration":"本页配图画面描述","design":{{"composition":"本页图文关系与理由","image_position":"right|left|background|full_bleed|inline|none","density":"airy|balanced|dense","emphasis":"第一眼应理解的判断"}},"question":"仅互动题","options":["选项"],"answer":0,"explanation":"答案讲解"}}]}}

要求：
- 每页 ONE KEY IDEA：一页只讲一个核心想法，不要堆砌五条并列 bullet；
- 每页都必须有 narration，内容与页面一致；
- 图文结合：除封面外每页输出 illustration 字段——一个 30 字以内、可直接画出来的教学画面描述
  （具体主体 + 空间关系，如「水杯中的筷子向上弯折，水面处错开」）；无合适画面的页写 ""；
- 页面结构完整、事实可靠、中文自然；
- 内容深度匹配学习者真实背景（初学者更直观、一步一个概念；进阶学习者少基础定义、多边界与迁移）；
- 输出前逐项核对用户原话的页数、逐页条件、明确排除项；不得用通用封面或总结页挤占用户指定的内容页。
- {interaction_rule}
- 只输出 JSON。"""
    from services.llm_service import get_llm_service
    llm = await get_llm_service()
    raw = await llm.generate(prompt, max_tokens=3600)
    deck = normalize_slide_deck(_extract_json(raw), include_interactions=include_interactions)
    if slide_count and len(deck["slides"]) != int(slide_count):
        correction = (
            f"上一次草稿生成了 {len(deck['slides'])} 页，但用户明确要求恰好 {slide_count} 页，"
            "封面与总结也计入总数。请重新组织完整内容，保留用户原话中的逐页要求；"
            "不要机械截断或附加空白页。只输出上述结构的严格 JSON。\n"
            f"原始任务与依据：{prompt[:7000]}"
        )
        raw = await llm.generate(correction, max_tokens=3600)
        deck = normalize_slide_deck(_extract_json(raw), include_interactions=include_interactions)
        if len(deck["slides"]) != int(slide_count):
            reason = f"要求 {slide_count} 页，当前初稿只有 {len(deck['slides'])} 页。"
            raise GenerationDraftBlocked(
                "PRESENTATION_SLIDE_COUNT_MISMATCH", reason=reason,
                body=presentation_draft(deck, reason=reason),
            )
    # 内容质量闸门：页数对了、结构合法了，不代表这一页页真的在讲这个主题。
    # 放在插图之前 —— 让一段跑题的内容先过闸，再谈配图，免得白烧图片接口。
    from services.content_quality_gate import (
        KIND_PRESENTATION,
        quality_receipt,
        run_content_quality_gate,
    )

    async def revise_deck(*, existing_candidate, constraints, prior_review, round_number):
        rewrite_prompt = (
            f"{prompt}\n\n上一稿未通过内容质量检查，请按下面的要求重做整份 PPT：\n"
            + "\n".join(f"- {item}" for item in constraints)
            + "\n\n上一稿：" + str(json.dumps(dict(existing_candidate), ensure_ascii=False))[:6000]
            + "\n保持同样的页数要求与输出结构，只输出 JSON。"
        )
        rewritten_raw = await llm.generate(rewrite_prompt, max_tokens=3600)
        return normalize_slide_deck(_extract_json(rewritten_raw), include_interactions=include_interactions)

    gate = await run_content_quality_gate(
        kind=KIND_PRESENTATION,
        candidate=deck,
        topic=topic,
        learning_goal=learning_goal,
        revise=revise_deck,
        max_rounds=3,
    )
    gated_deck = gate["candidate"]
    if slide_count and len(gated_deck.get("slides") or []) != int(slide_count):
        # 改写不能把页数要求改掉；宁可留着带问题的原稿，也不能悄悄少给几页。
        gated_deck = deck
    deck = gated_deck

    deck["provider"] = str(getattr(settings, "LLM_PROVIDER", "unknown") or "unknown").lower()
    if deck["provider"] == "mock":
        raise RuntimeError("presentation generation requires a real LLM provider")
    # 版本级主题快照：生成时定格，之后不随产品主题变化。
    deck["presentation_theme"] = theme_key
    deck["content_review"] = quality_receipt(gate)
    # 这份 PPT 的每一条事实应该能追到来源。searched=false 时如实为空 ——
    # 一份没有来源的作品不该看起来像有来源。
    deck["grounding"] = {
        "searched": bool(grounding.get("searched")),
        "source_refs": list(grounding.get("source_refs") or []),
        "note": str(grounding.get("note") or ""),
        "sufficient": bool(grounding.get("sufficient")) if grounding.get("searched") else False,
    }
    # 封面插图技能：为首页配一张与主题匹配、不承载文字的示意图；失败静默跳过。
    # 模型把封面判成 image_position=none 时直接不生成——渲染器本来就不会画它，
    # 照出图只是每次生成白烧一次图片接口。
    cover = (deck.get("slides") or [{}])[0]
    if str(_mapping(cover.get("design")).get("image_position") or "").strip().lower() != "none":
        try:
            from services.illustration_skill import generate_material_illustration

            outcome = await generate_material_illustration(
                f"{topic} 的核心概念直观场景",
                theme=theme_key,
                style_hint=(presentation_style or "")[:40],
            )
            if outcome.get("ok") and outcome.get("url"):
                cover["image"] = outcome["url"]
        except Exception:
            pass
    # 图文结合：为其余每页生成一张与该页内容匹配的教学插图（并发有界，逐页静默降级）。
    await _attach_slide_illustrations(deck, theme_key=theme_key, presentation_style=presentation_style)
    return deck


async def _attach_slide_illustrations(
    deck: dict, *, theme_key: str, presentation_style: str | None,
) -> int:
    """Give every slide without an image a content-matched teaching illustration.

    The per-slide illustration concept comes from the model's own `illustration`
    field when present, else from the slide title plus its first visible line.
    Individual failures degrade that slide to text-only; the deck always ships.

    Two kinds of page are deliberately left without a picture, because both the
    art-direction and the presentation-authoring skill say so: a page whose
    `design.image_position` is `none` (the model decided text carries it better)
    and a quiz page that declared no picture concept (the skill's own example of
    a page that must not have one bolted on).
    """
    import asyncio

    from services.illustration_skill import generate_material_illustration

    def _wants_picture(slide: dict) -> bool:
        design = _mapping(slide.get("design"))
        if str(design.get("image_position") or "").strip().lower() == "none":
            return False
        if slide.get("interaction") and not _text(slide.get("illustration"), 120):
            return False
        return True

    slides = [s for s in (deck.get("slides") or []) if isinstance(s, dict)]
    targets = [
        s for s in slides
        if not str(s.get("image") or "").strip() and _wants_picture(s)
    ]
    if not targets:
        return 0
    semaphore = asyncio.Semaphore(3)

    async def _one(slide: dict) -> bool:
        concept = _text(slide.get("illustration"), 120)
        if len(concept) < 8:
            first_item = next((i for i in (slide.get("items") or []) if _text(i, 30)), "")
            concept = f"{_text(slide.get('title'), 60)}：{_text(first_item, 60)}".strip("：")
        if len(concept) < 8:
            return False
        try:
            async with semaphore:
                outcome = await generate_material_illustration(
                    concept, theme=theme_key,
                    style_hint=(presentation_style or "")[:40],
                )
        except Exception:
            return False
        if outcome.get("ok") and outcome.get("url"):
            slide["image"] = str(outcome["url"])
            return True
        return False

    results = await asyncio.gather(*(_one(s) for s in targets))
    return sum(1 for ok in results if ok)


async def revise_workspace_presentation(
    *, existing_deck: Mapping[str, Any], instruction: str, delivery_mode: str,
) -> dict:
    """Revise one persisted deck through the real provider, keeping its delivery contract."""
    include_interactions = delivery_mode == "interactive_lesson"
    existing_json = json.dumps(dict(existing_deck), ensure_ascii=False)[:18000]
    prompt = f"""{skill_prompt_prefix("resource-art-direction", "presentation-authoring", max_chars_per_skill=1600)}你是 PPT 与互动课堂内容修改智能体。请根据修改要求更新现有产物。
只输出完整、严格的 JSON，不要输出 Markdown 代码围栏，也不要只输出变更片段。
交付模式：{delivery_mode}
修改要求：{_text(instruction, 1000)}
现有产物：{existing_json}

输出结构：
{{"title":"标题","slides":[{{"type":"title|goals|content|code|quiz|summary","title":"页标题","subtitle":"可选","body":"可选正文","items":["要点"],"code":"可选代码","lang":"可选语言","narration":"本页自然讲解","illustration":"本页配图画面描述","design":{{"composition":"本页图文关系与理由","image_position":"right|left|background|full_bleed|inline|none","density":"airy|balanced|dense","emphasis":"第一眼应理解的判断"}},"question":"仅互动题","options":["选项"],"answer":0,"explanation":"答案讲解"}}]}}

要求：
- 只执行用户要求的修改，保留未要求改变的主题与有效内容；
- 保持现有页数和页面顺序不变，不新增、删除或移动页面；
- 如果修改要求指定“第 N 页”或“page N”，必须直接改写 slides 数组中对应位置的页面，不能在末尾追加替代页；
- 目标页的 title、subtitle、body、items 或 narration 至少有一个可见字段必须实际改变，不能原样返回；
- 只有目标页的 design 允许随内容调整；未被要求的页面必须原样保留它的 illustration 与 design，
  不得把配图位置、疏密或强调句重置成默认值；
- 内容改到图文关系也变化时，同步改写该页 design.composition 与 emphasis，而不是只改文字；
- 不得发明来源资料中没有的指标、用户结果、研究、引文、证据或产品能力；缺少证据时用需要展示或核查的要求表达；
- 每页都保留或补全 narration；
- 互动课堂保留可作答结构，普通 PPT 不包含答案元数据；
- 只输出 JSON。"""
    from services.llm_service import get_llm_service
    llm = await get_llm_service()
    raw = await llm.generate(prompt, max_tokens=3600, response_format="json")
    deck = normalize_slide_deck(_extract_json(raw), include_interactions=include_interactions)
    deck["provider"] = str(getattr(settings, "LLM_PROVIDER", "unknown") or "unknown").lower()
    if deck["provider"] == "mock":
        raise RuntimeError("presentation revision requires a real LLM provider")
    # 学习者自己要求的这次修订同样要过内容质量闸门：修订能把一页改错，而以前
    # 这条路径不做任何内容检查就发布新版本。
    from services.content_quality_gate import (
        KIND_PRESENTATION, quality_receipt, run_content_quality_gate,
    )

    async def revise_revised_deck(*, existing_candidate, constraints, prior_review, round_number):
        rewrite_prompt = (
            f"{prompt}\n\n上一稿未通过内容质量检查，请按下面的要求重做：\n"
            + "\n".join(f"- {item}" for item in constraints)
            + "\n\n上一稿：" + str(json.dumps(dict(existing_candidate), ensure_ascii=False))[:6000]
            + "\n只输出 JSON。"
        )
        return normalize_slide_deck(
            _extract_json(await llm.generate(rewrite_prompt, max_tokens=3600)),
            include_interactions=include_interactions,
        )

    revision_gate = await run_content_quality_gate(
        kind=KIND_PRESENTATION, candidate=deck,
        topic=_text(existing_deck.get("topic") or existing_deck.get("title"), 200),
        revise=revise_revised_deck,
    )
    deck = dict(revision_gate["candidate"])
    deck["content_review"] = quality_receipt(revision_gate)
    # 主题快照是版本属性：修订不改变它。
    deck["presentation_theme"] = normalize_presentation_theme(existing_deck.get("presentation_theme"))
    # 图文结合：修订重建了页面，给丢失配图的页重新生成插图。
    await _attach_slide_illustrations(
        deck, theme_key=str(deck["presentation_theme"]),
        presentation_style=_text(existing_deck.get("presentation_style"), 40) or None,
    )
    return deck


def _text_width_units(text: str) -> int:
    # 近似排版宽度（半角单位）：CJK/全角字符按 2 个单位计，其余按 1。
    return sum(2 if ord(ch) >= 0x2E80 else 1 for ch in text)


def _estimate_wrapped_lines(text: str, *, units_per_line: int) -> int:
    lines = 0
    for raw in str(text or "").split("\n"):
        units = _text_width_units(raw.rstrip())
        lines += max(1, -(-units // max(1, units_per_line)))
    return max(1, lines)


def _slide_text(slide: Mapping[str, Any]) -> str:
    parts = []
    if slide.get("subtitle"):
        parts.append(_text(slide["subtitle"], 500))
    if slide.get("body"):
        parts.append(_text(slide["body"], 4000))
    parts.extend(f"• {_text(item, 500)}" for item in (slide.get("items") or []))
    if slide.get("code"):
        parts.append(_text(slide["code"], 5000))
    interaction = slide.get("interaction") or {}
    if interaction:
        parts.append(_text(interaction.get("question"), 500))
        parts.extend(f"{chr(65 + i)}. {_text(option, 300)}" for i, option in enumerate(interaction.get("options") or []))
    return "\n".join(part for part in parts if part)


# ---------------------------------------------------------------------------
# Morandi PPTX renderer：一个渲染引擎，两套 theme token，全部可编辑
# （文字 = text frame；图形 = autoshape；代码 = 文本块；不做整页贴图）
# ---------------------------------------------------------------------------
# 版式密度：由每页 design.density 决定（技能与模型给出），渲染器只负责忠实实现。
# max_items 是密度真正的语义——airy 页要留白给画面，不是把同样多的字排松一点；
# presentation-authoring 规定正文不超过 5 条，所以 dense 也止于 5。
# balanced 一档刻意等于历史几何，让没有 design 的旧产物保持原样。
_DENSITY_METRICS: dict[str, dict[str, float]] = {
    "airy": {
        "max_items": 3, "step": 1.02, "item_size": 20, "body_size": 19, "body_ls": 1.5,
        "emphasis_size": 23, "subtitle_size": 20, "measure": 0.78,
    },
    "balanced": {
        "max_items": 5, "step": 0.82, "item_size": 18, "body_size": 18, "body_ls": 1.35,
        "emphasis_size": 21, "subtitle_size": 20, "measure": 1.0,
    },
    "dense": {
        "max_items": 5, "step": 0.66, "item_size": 16, "body_size": 16, "body_ls": 1.2,
        "emphasis_size": 18, "subtitle_size": 17, "measure": 1.0,
    },
}

_IMAGE_POSITIONS = {"right", "left", "background", "full_bleed", "inline", "none"}


# ---------------------------------------------------------------------------
# 版式计划（layout plan）——几何决策的唯一权威
#
# PPTX 导出与 PNG 栅格图都消费同一份计划，所以"用户打开的文件""页面上的预览"
# "智能体用来审自己的那张图"不可能再各画各的。计划只描述图形与几何，不含任何
# 渲染库类型；坐标一律是 EMU（1/914400 英寸），颜色一律是主题令牌的键名。
# ---------------------------------------------------------------------------
EMU_PER_IN = 914400
SLIDE_W_EMU = int(round(13.333 * EMU_PER_IN))
SLIDE_H_EMU = int(round(7.5 * EMU_PER_IN))


def _emu_inches(value: float) -> int:
    return int(round(float(value) * EMU_PER_IN))


def _pt(points: float) -> int:
    return int(round(float(points) * 12700))


def resolve_static_picture(url: Any) -> Path | None:
    """Only product-owned static illustrations may be embedded or rasterised."""
    candidate = str(url or "").strip()
    if not candidate.startswith("/static/"):
        return None
    path = Path(__file__).resolve().parent.parent / candidate.lstrip("/")
    return path if path.is_file() else None


def picture_source_size(path: Path) -> tuple[int, int] | None:
    try:
        from PIL import Image

        with Image.open(path) as probe:
            return probe.size
    except Exception:
        return None


def cover_crop(path: Path, box_w: int, box_h: int):
    """Crop a picture to the box aspect so a bleed page really bleeds."""
    from PIL import Image

    with Image.open(path) as probe:
        image = probe.convert("RGB")
    source_w, source_h = image.size
    if source_w <= 0 or source_h <= 0 or box_w <= 0 or box_h <= 0:
        return image
    ratio = box_w / box_h
    if source_w / source_h > ratio:
        crop_w = max(1, int(source_h * ratio))
        left = (source_w - crop_w) // 2
        return image.crop((left, 0, left + crop_w, source_h))
    crop_h = max(1, int(source_w / ratio))
    top = (source_h - crop_h) // 2
    return image.crop((0, top, source_w, top + crop_h))


def build_slide_plan(deck: Mapping[str, Any]) -> dict:
    """Turn one deck into a renderer-independent layout plan."""
    theme_key = normalize_presentation_theme(deck.get("presentation_theme"))
    shapes: list[dict] = []

    def rect(x, y, w, h, color, *, round_shape=False, radius=None, line_color=None):
        box = (int(x), int(y), int(w), int(h))
        shapes.append({
            "kind": "rect", "box": box, "color": str(color),
            "round": bool(round_shape), "radius": radius,
            "line": str(line_color) if line_color else None,
        })
        return box

    def text(x, y, w, h, value, *, size=18, color="text", bold=False, font="body",
             align="left", anchor="top", line_spacing=1.15):
        box = (int(x), int(y), int(w), int(h))
        shapes.append({
            "kind": "text", "box": box, "text": str(value or ""),
            "size": float(size), "color": str(color), "bold": bool(bold),
            "font": str(font), "align": str(align), "anchor": str(anchor),
            "line_spacing": float(line_spacing),
        })
        return box

    def picture(url, x, y, w, h, *, cover=False):
        path = resolve_static_picture(url)
        if path is None:
            return None
        box_w, box_h = int(w), int(h)
        if not cover:
            size = picture_source_size(path)
            if size and size[0] > 0 and size[1] > 0:
                scale = min(box_w / size[0], box_h / size[1])
                box_w, box_h = int(size[0] * scale), int(size[1] * scale)
        box_x = int(x) + (int(w) - box_w) // 2
        box_y = int(y) + (int(h) - box_h) // 2
        shapes.append({
            "kind": "picture", "box": (box_x, box_y, box_w, box_h),
            "path": str(path), "cover": bool(cover),
        })
        return (box_x, box_y, box_w, box_h)

    def theme_motif():
        # 轻量主题图形语言：Kaki = 圆角纸片 + 一颗柿子色圆点；Ocean = 顶部细线 + 小方块。
        if theme_key == "kaki":
            rect(SLIDE_W_EMU - _emu_inches(1.5), _emu_inches(0.42),
                 _emu_inches(0.26), _emu_inches(0.26), "primary", round_shape=True, radius=0.5)
        else:
            rect(0, 0, SLIDE_W_EMU, _pt(3), "primary")
            rect(SLIDE_W_EMU - _emu_inches(1.5), _emu_inches(0.42),
                 _emu_inches(0.26), _emu_inches(0.26), "secondary")

    slides = deck.get("slides") or []
    total = len(slides)
    pages: list[dict] = []

    for index, slide_data in enumerate(slides):
        shapes = []
        rect(0, 0, SLIDE_W_EMU, SLIDE_H_EMU, "background")
        theme_motif()

        kind = str(slide_data.get("type") or "content")
        title = _text(slide_data.get("title"), 200)
        subtitle = _text(slide_data.get("subtitle"), 300)
        body = _text(slide_data.get("body"), 4000)
        items = [str(item) for item in (slide_data.get("items") or []) if str(item).strip()]
        code = _text(slide_data.get("code"), 5000)
        interaction = slide_data.get("interaction") or {}
        design = _mapping(slide_data.get("design"))
        image_position = str(design.get("image_position") or "inline").lower()
        if image_position not in _IMAGE_POSITIONS:
            image_position = "inline"
        density = str(design.get("density") or "balanced").lower()
        if density not in _DENSITY_METRICS:
            density = "balanced"
        metrics = _DENSITY_METRICS[density]
        emphasis = _text(design.get("emphasis"), 180)
        image_url = _text(slide_data.get("image"), 600)
        is_cover = index == 0 or kind == "title"

        # 通版图铺满画布，再压一层纸面面板：标题与正文始终可读，与网页预览同一策略。
        panel_right: int | None = None
        if image_url and image_position in {"background", "full_bleed"}:
            picture(image_url, 0, 0, SLIDE_W, SLIDE_H, cover=True)
            panel_w = _emu_inches(9.4) if image_position == "full_bleed" else _emu_inches(10.7)
            rect(_emu_inches(0.7), _emu_inches(0.5), panel_w, _emu_inches(6.4), "surface")
            panel_right = int(_emu_inches(0.7)) + int(panel_w)

        # 标题区（32–40pt，克制）：全篇统一，让阅读节奏稳定；页面差异由正文区承担。
        title_size = 40 if is_cover else 32
        title_w = int(int(panel_right) - int(_emu_inches(1.53))) if panel_right else _emu_inches(10.9)
        rect(_emu_inches(0.9), _emu_inches(0.82), _emu_inches(0.09), _emu_inches(0.62), "primary")
        text(_emu_inches(1.18), _emu_inches(0.72), title_w, _emu_inches(1.0),
                 title or (deck.get("title") or ""), size=title_size, bold=True, anchor="middle")

        if is_cover:
            cover_top = 2.35
            if emphasis:
                text(_emu_inches(1.2), _emu_inches(cover_top), title_w, _emu_inches(0.66),
                         emphasis, size=metrics["emphasis_size"], color="primary", bold=True)
                cover_top += 0.74
            if subtitle:
                text(_emu_inches(1.2), _emu_inches(cover_top), title_w, _emu_inches(0.8),
                         subtitle, size=22, color="muted")
            meta = _text(deck.get("title"), 120)
            if meta:
                text(_emu_inches(1.2), _emu_inches(6.35), _emu_inches(8.0), _emu_inches(0.5),
                         meta, size=13, color="muted")
        elif kind == "goals":
            goals = items or ([body] if body else [])
            goals_top = 2.15
            if emphasis:
                text(_emu_inches(1.2), _emu_inches(1.76), _emu_inches(10.9), _emu_inches(0.52),
                         emphasis, size=metrics["emphasis_size"], color="primary", bold=True)
                goals_top = 2.42
            step = max(0.72, metrics["step"] * 1.16)
            goal_h = max(0.5, step - 0.08)
            fit = 1 + int((6.92 - goal_h - goals_top) / step)
            for gi, goal in enumerate(goals[: max(1, min(len(goals), fit))]):
                y = _emu_inches(goals_top + gi * step)
                rect(_emu_inches(1.2), y, _emu_inches(0.42), _emu_inches(0.42), "secondary",
                         round_shape=True, radius=0.5)
                text(_emu_inches(1.2), y, _emu_inches(0.42), _emu_inches(0.42),
                         str(gi + 1), size=16, color="background", bold=True,
                         align="center", anchor="middle")
                text(_emu_inches(1.85), y - _emu_inches(0.04), _emu_inches(9.9), _emu_inches(goal_h),
                         goal, size=metrics["item_size"] + 2)
        elif kind == "code" or code:
            if subtitle:
                text(_emu_inches(1.2), _emu_inches(1.75), _emu_inches(10.6), _emu_inches(0.6),
                         subtitle, size=18, color="muted")
            # 代码页布局预算（两套主题共享同一几何）：
            #   code_bottom + 0.15in <= note_top，note_bottom + 0.14in <= footer_top(7.06in)。
            # 面板高度随代码行数自适应；代码过长时按档位缩小字号（最低 11pt）。
            # 行高按 PowerPoint 实测校准：倍数行距作用在字体自然行高上
            # （Consolas 14pt×1.25 实测 ≈1.50em，含 CJK 注释行 ≈1.65em；雅黑正文 ≈1.41em）。
            note = _text(slide_data.get("narration"), 150)
            note_h = 0.0
            if note:
                note_lines = _estimate_wrapped_lines(note, units_per_line=119)
                note_h = max(0.36, note_lines * 11.5 * 1.45 / 72 + 0.12)
            max_panel_bottom = 6.92 - (note_h + 0.15 if note else 0.0)
            panel_x, panel_y = _emu_inches(1.2), _emu_inches(2.3)
            panel_w = _emu_inches(10.9)
            code_rows = (code or "").split("\n")
            code_size = 11
            panel_h = max(1.5, max_panel_bottom - 2.3)
            code_columns: list[str] | None = None
            for candidate_size in (14, 13, 12, 11):
                code_line_h = candidate_size * 1.65 / 72
                code_units_per_line = int((10.2 * 72 - 14.4) / (candidate_size * 0.55) * 0.9)
                code_lines = _estimate_wrapped_lines(code, units_per_line=code_units_per_line)
                needed_h = max(1.5, code_lines * code_line_h + 0.5)
                if 2.3 + needed_h <= max_panel_bottom:
                    code_size, panel_h = candidate_size, needed_h
                    break
            else:
                # 单列放不下：保持 11pt 下限，代码等分成两列（优先在空行处切分）。
                code_size = 11
                col_line_h = code_size * 1.65 / 72
                col_units_per_line = int(((10.9 - 1.2) / 2 * 72 - 14.4) / (code_size * 0.55) * 0.9)

                def _col_lines(rows: list[str]) -> int:
                    return _estimate_wrapped_lines("\n".join(rows), units_per_line=col_units_per_line)

                n = len(code_rows)
                lo, hi = max(1, int(n * 0.35)), min(n - 1, int(n * 0.65))
                blanks = [i for i in range(lo, hi + 1) if not code_rows[i].strip()]
                split = min(blanks, key=lambda i: abs(i - n / 2)) if blanks else n // 2
                top_rows, bottom_rows = code_rows[:split], code_rows[split:]
                panel_h = max(
                    1.5,
                    max(_col_lines(top_rows), _col_lines(bottom_rows)) * col_line_h + 0.5,
                )
                if 2.3 + panel_h > max_panel_bottom:
                    panel_h = max(1.5, max_panel_bottom - 2.3)
                code_columns = ["\n".join(top_rows), "\n".join(bottom_rows)]
            panel_h_emu = _emu_inches(panel_h)
            rect(panel_x, panel_y, panel_w, panel_h_emu, "code_background",
                     round_shape=True, radius=0.06)
            if code_columns:
                col_w = _emu_inches((10.9 - 1.2) / 2)
                for ci, column_text in enumerate(code_columns):
                    text(panel_x + _emu_inches(0.35) + ci * (col_w + _emu_inches(0.5)),
                             panel_y + _emu_inches(0.28), col_w, panel_h_emu - _emu_inches(0.5),
                             column_text, size=code_size, color="code_text",
                             font="code", line_spacing=1.25)
            else:
                text(panel_x + _emu_inches(0.35), panel_y + _emu_inches(0.28),
                         panel_w - _emu_inches(0.7), panel_h_emu - _emu_inches(0.5),
                         code or "", size=code_size, color="code_text", font="code", line_spacing=1.25)
            if note:
                text(_emu_inches(1.2), panel_y + panel_h_emu + _emu_inches(0.15),
                         _emu_inches(10.6), _emu_inches(note_h),
                         note, size=11.5, color="muted")
        elif interaction:
            text(_emu_inches(1.2), _emu_inches(1.85), _emu_inches(10.9), _emu_inches(1.0),
                     interaction.get("question") or "", size=24, bold=True)
            options = [str(o) for o in (interaction.get("options") or [])][:4]
            for oi, option in enumerate(options):
                col, row = oi % 2, oi // 2
                x = _emu_inches(1.2 + col * 5.55)
                y = _emu_inches(3.3 + row * 1.15)
                rect(x, y, _emu_inches(5.25), _emu_inches(0.92), "surface",
                         round_shape=True, radius=0.18, line_color="line")
                text(x + _emu_inches(0.25), y, _emu_inches(4.8), _emu_inches(0.92),
                         f"{chr(65 + oi)}. {option}", size=16, anchor="middle")
            text(_emu_inches(1.2), _emu_inches(6.05), _emu_inches(10.6), _emu_inches(0.5),
                     "想一想再作答；作答会记入学习安排。", size=13, color="muted")
        else:
            # 正文区版式：图文关系取自 image_position，疏密取自 density，
            # "第一眼该看见什么" 取自 emphasis。三者都由技能与模型按内容决定，
            # 渲染器只把它们变成几何，不再有一套所有页面共用的固定排版。
            zone_x, zone_w = 1.2, 10.9
            if panel_right is not None:
                zone_w = (panel_right - int(_emu_inches(1.2)) - int(_emu_inches(0.5))) / EMU_PER_IN
            zone_top = 1.85

            # 左右并置先定列、再排文字：文字栏定下来以后，强调句就只能待在文字栏内，
            # 不会排成横跨全宽的一行伸到图下面（那样只有靠图片适配框恰好更低才不重叠）。
            if image_url and image_position in {"left", "right"}:
                picture_x = 7.45 if image_position == "right" else 1.2
                placed = picture(image_url, _emu_inches(picture_x), _emu_inches(1.85),
                                               _emu_inches(4.65), _emu_inches(4.6))
                if placed is not None:
                    zone_x = 1.2 if image_position == "right" else 6.25
                    zone_w = 5.3

            if emphasis:
                text(_emu_inches(zone_x), _emu_inches(zone_top - 0.07), _emu_inches(zone_w), _emu_inches(0.55),
                         emphasis, size=metrics["emphasis_size"], color="primary", bold=True)
                zone_top += 0.7

            if image_url and image_position == "inline":
                # 内嵌图先给自己定一个上限：这一页的正文不能被图挤没，
                # 至少留出两条的位置，剩下的高度才归画面。
                text_floor = 2 * min(metrics["step"], 0.82) if items else 0.8
                cap = max(1.4, 6.92 - zone_top - (0.62 if subtitle else 0.0) - text_floor)
                placed = picture(image_url, _emu_inches(zone_x), _emu_inches(zone_top),
                                               _emu_inches(zone_w), _emu_inches(cap))
                if placed is not None:
                    zone_top = (placed[1] + placed[3]) / EMU_PER_IN + 0.24

            measure = zone_w * metrics["measure"]
            if subtitle:
                text(_emu_inches(zone_x), _emu_inches(zone_top - 0.07), _emu_inches(measure), _emu_inches(0.62),
                         subtitle, size=metrics["subtitle_size"], color="primary")
                zone_top += 0.62

            # summary 的收束条只在页面还有空间时出现，避免和正文抢位置。
            band = 1.05 if (kind == "summary" and zone_top <= 5.5) else 0.0
            if items:
                limit = min(len(items), int(metrics["max_items"]))
                step = metrics["step"]
                available = 6.92 - band - zone_top
                if limit * step > available:
                    step = max(0.55, available / limit)
                # 密度定上限，版面定实际条数：放不下就少放，绝不越过页脚。
                limit = min(limit, max(0, int(available // step)))
                for ii, item in enumerate(items[:limit]):
                    y = zone_top + ii * step
                    rect(_emu_inches(zone_x + 0.05), _emu_inches(y + 0.12), _emu_inches(0.16), _emu_inches(0.16),
                             "accent", round_shape=True, radius=0.5)
                    text(_emu_inches(zone_x + 0.42), _emu_inches(y), _emu_inches(measure - 0.42),
                             _emu_inches(min(0.75, step)), item, size=metrics["item_size"])
            elif body and 6.92 - band - zone_top >= 0.6:
                text(_emu_inches(zone_x), _emu_inches(zone_top - 0.07), _emu_inches(measure),
                         _emu_inches(6.92 - band - zone_top), body,
                         size=metrics["body_size"], line_spacing=metrics["body_ls"])
            if band:
                closing = items[-1] if items else (body or "").split("\n")[0]
                if closing:
                    rect(_emu_inches(1.2), _emu_inches(6.0), _emu_inches(10.9), _emu_inches(0.85), "surface",
                             round_shape=True, radius=0.25, line_color="line")
                    text(_emu_inches(1.5), _emu_inches(6.0), _emu_inches(10.3), _emu_inches(0.85),
                             closing, size=17, color="accent", bold=True, anchor="middle")

        # 页脚：页码 + 细线（低存在感）
        rect(_emu_inches(1.2), _emu_inches(7.06), _emu_inches(10.9), _pt(1), "line")
        text(_emu_inches(11.6), _emu_inches(7.12), _emu_inches(0.9), _emu_inches(0.34),
                 f"{index + 1} / {total}", size=12, color="muted", align="right")

        notes = _text(slide_data.get("narration"), 5000)
        pages.append({"shapes": shapes, "notes": notes})

    return {"theme": theme_key, "pages": pages}


def build_pptx_bytes(deck: Mapping[str, Any]) -> bytes:
    """Render the layout plan into a valid, fully editable PPTX package."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    plan = build_slide_plan(deck)
    theme = presentation_theme_tokens(plan["theme"])
    body_font = "Microsoft YaHei"
    code_font = "Consolas"
    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    anchor_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}

    def rgb(key: str) -> RGBColor:
        return RGBColor.from_string(theme[key])

    presentation = Presentation()
    presentation.slide_width = Emu(SLIDE_W_EMU)
    presentation.slide_height = Emu(SLIDE_H_EMU)
    blank = presentation.slide_layouts[6]

    for page in plan["pages"]:
        slide = presentation.slides.add_slide(blank)
        for shape in page["shapes"]:
            box = shape["box"]
            if shape["kind"] == "rect":
                element = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE if shape["round"] else MSO_SHAPE.RECTANGLE,
                    Emu(box[0]), Emu(box[1]), Emu(box[2]), Emu(box[3]),
                )
                element.fill.solid()
                element.fill.fore_color.rgb = rgb(shape["color"])
                if shape["line"]:
                    element.line.color.rgb = rgb(shape["line"])
                    element.line.width = Pt(1)
                else:
                    element.line.fill.background()
                element.shadow.inherit = False
                if shape["round"] and shape["radius"] is not None:
                    try:
                        element.adjustments[0] = shape["radius"]
                    except Exception:
                        pass
            elif shape["kind"] == "text":
                element = slide.shapes.add_textbox(
                    Emu(box[0]), Emu(box[1]), Emu(box[2]), Emu(box[3]),
                )
                frame = element.text_frame
                frame.word_wrap = True
                frame.vertical_anchor = anchor_map.get(shape["anchor"], MSO_ANCHOR.TOP)
                for i, line in enumerate(shape["text"].split("\n")):
                    paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    paragraph.alignment = align_map.get(shape["align"], PP_ALIGN.LEFT)
                    paragraph.line_spacing = shape["line_spacing"]
                    run = paragraph.add_run()
                    run.text = line
                    run.font.size = Pt(shape["size"])
                    run.font.bold = shape["bold"]
                    run.font.name = code_font if shape["font"] == "code" else body_font
                    run.font.color.rgb = rgb(shape["color"])
            elif shape["kind"] == "picture":
                source: Any = shape["path"]
                if shape["cover"]:
                    buffer = io.BytesIO()
                    cover_crop(Path(shape["path"]), box[2], box[3]).save(buffer, format="PNG")
                    buffer.seek(0)
                    source = buffer
                try:
                    slide.shapes.add_picture(source, Emu(box[0]), Emu(box[1]), Emu(box[2]), Emu(box[3]))
                except Exception:
                    continue
        if page.get("notes"):
            try:
                slide.notes_slide.notes_text_frame.text = page["notes"]
            except Exception:
                pass

    payload = io.BytesIO()
    presentation.save(payload)
    return payload.getvalue()
