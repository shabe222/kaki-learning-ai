"""Real in-memory PDF and DOCX exports for persisted workspace lectures."""
from io import BytesIO
import os
import re
import textwrap


def _plain_lines(markdown: str):
    for raw in str(markdown or "").splitlines():
        line = re.sub(r"^\s{0,3}#{1,6}\s*", "", raw)
        line = re.sub(r"^\s*[-*+]\s+", "• ", line)
        line = re.sub(r"[*_`]+", "", line).strip()
        if line:
            yield line


def build_pdf_bytes(title: str, markdown: str) -> bytes:
    """Build a valid, readable multipage PDF using Pillow and a CJK font."""
    from PIL import Image, ImageDraw, ImageFont

    font_path = next((path for path in (
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ) if os.path.exists(path)), None)
    if not font_path:
        raise RuntimeError("no usable export font")
    title_font = ImageFont.truetype(font_path, 34)
    body_font = ImageFont.truetype(font_path, 22)
    width, height, margin = 1240, 1754, 90
    pages = []
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    y = margin

    def new_page():
        nonlocal image, draw, y
        pages.append(image)
        image = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(image)
        y = margin

    draw.text((margin, y), str(title or "学习讲义"), fill="#24211d", font=title_font)
    y += 70
    for line in _plain_lines(markdown):
        for wrapped in textwrap.wrap(line, width=42, replace_whitespace=False) or [""]:
            if y > height - margin - 40:
                new_page()
            draw.text((margin, y), wrapped, fill="#33302c", font=body_font)
            y += 38
        y += 10
    pages.append(image)
    output = BytesIO()
    pages[0].save(output, "PDF", resolution=120.0, save_all=True, append_images=pages[1:])
    return output.getvalue()


def build_docx_bytes(title: str, markdown: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(str(title or "学习讲义"), 0)
    for raw in str(markdown or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
        elif re.match(r"^[-*+]\s+", line):
            document.add_paragraph(re.sub(r"^[-*+]\s+", "", line), style="List Bullet")
        else:
            document.add_paragraph(re.sub(r"[*_`]+", "", line))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
